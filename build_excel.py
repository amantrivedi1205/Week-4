"""Builds excel/Retail_Sales_Analysis.xlsx: cleaned data plus formula-driven summary tables and charts."""
import json
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
df = pd.read_csv(ROOT / "data" / "retail_sales_clean.csv", parse_dates=["Order Date"])
metrics = json.load(open(ROOT / "reports" / "metrics.json"))
N = len(df)
LAST = N + 1  # last data row

F = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
INPUT = Font(name=F, size=10, color="0000FF")
TITLE = Font(name=F, size=14, bold=True, color="1F3A5F")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
USD, PCT, INT = '$#,##0;($#,##0);"-"', '0.0%;(0.0%);"-"', '#,##0'

wb = Workbook()

# ---------------- Clean_Data ----------------
ws = wb.active; ws.title = "Clean_Data"
cols = ["Order ID", "Order Date", "Ship Mode", "Segment", "Region", "Category", "Sub-Category",
        "Quantity", "Discount", "Sales", "Profit"]
ws.append(cols + ["Year", "Profit Margin"])
for r in df[cols].itertuples(index=False):
    ws.append([r[0], r[1].to_pydatetime(), *r[2:]])
for i in range(2, LAST + 1):
    ws[f"L{i}"] = f"=YEAR(B{i})"
    ws[f"M{i}"] = f"=IF(J{i}>0,K{i}/J{i},0)"
    ws[f"B{i}"].number_format = "yyyy-mm-dd"
    ws[f"I{i}"].number_format = "0%"
    ws[f"J{i}"].number_format = '$#,##0.00'; ws[f"K{i}"].number_format = '$#,##0.00;($#,##0.00)'
    ws[f"M{i}"].number_format = "0.0%;(0.0%)"
for c in ws[1]:
    c.fill, c.font, c.alignment = HDR_FILL, HDR_FONT, Alignment(horizontal="center")
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = BODY
for i, w in enumerate([16, 12, 15, 13, 10, 16, 14, 9, 9, 12, 12, 7, 12], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:M{LAST}"

rng = lambda col: f"Clean_Data!${col}$2:${col}${LAST}"

# ---------------- Summary (KPIs + tables) ----------------
s = wb.create_sheet("Summary", 0)
s["A1"] = "Retail Sales Analysis: Summary"; s["A1"].font = TITLE
s["A2"] = "All figures are live formulas over the Clean_Data sheet."; s["A2"].font = Font(name=F, size=9, italic=True)

def header(row, labels, col=1):
    for j, t in enumerate(labels):
        c = s.cell(row=row, column=col + j, value=t)
        c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, Alignment(horizontal="center", wrap_text=True), BORDER

def put(row, col, val, fmt=None, bold=False):
    c = s.cell(row=row, column=col, value=val); c.font = BOLD if bold else BODY; c.border = BORDER
    if fmt: c.number_format = fmt
    return c

# KPIs
header(4, ["KPI", "Value"])
kpis = [
    ("Total Sales", f"=SUM({rng('J')})", USD),
    ("Total Profit", f"=SUM({rng('K')})", USD),
    ("Profit Margin", "=B6/B5", PCT),
    ("Orders", f"=COUNTA({rng('A')})", INT),
    ("Average Order Value", "=B5/B8", USD),
    ("Average Discount", f"=AVERAGE({rng('I')})", PCT),
    ("Loss-Making Orders %", f'=COUNTIF({rng("K")},"<0")/B8', PCT),
]
for i, (k, f, fmt) in enumerate(kpis):
    put(5 + i, 1, k, bold=True); put(5 + i, 2, f, fmt)

def table(start, title, dim_col, items, extra_year=False):
    s.cell(row=start, column=1, value=title).font = Font(name=F, size=11, bold=True, color="1F3A5F")
    header(start + 1, [title.replace("By ", ""), "Sales", "Profit", "Margin"])
    for i, it in enumerate(items):
        r = start + 2 + i
        put(r, 1, it)
        put(r, 2, f"=SUMIFS({rng('J')},{rng(dim_col)},A{r})", USD)
        put(r, 3, f"=SUMIFS({rng('K')},{rng(dim_col)},A{r})", USD)
        put(r, 4, f"=IF(B{r}=0,0,C{r}/B{r})", PCT)
    end = start + 1 + len(items)
    r = end + 1
    put(r, 1, "Total", bold=True)
    put(r, 2, f"=SUM(B{start+2}:B{end})", USD, True); put(r, 3, f"=SUM(C{start+2}:C{end})", USD, True)
    put(r, 4, f"=IF(B{r}=0,0,C{r}/B{r})", PCT, True)
    return start + 2, end

cat_rows = table(14, "By Category", "F", sorted(df["Category"].unique()))
sub_items = (df.groupby("Sub-Category")["Profit"].sum().sort_values().index.tolist())
sub_start = cat_rows[1] + 4
sub_rows = table(sub_start, "By Sub-Category", "G", sub_items)
reg_start = sub_rows[1] + 4
reg_rows = table(reg_start, "By Region", "E", sorted(df["Region"].unique()))
seg_start = reg_rows[1] + 4
seg_rows = table(seg_start, "By Segment", "D", sorted(df["Segment"].unique()))

# Year table with YoY growth
yr_start = seg_rows[1] + 4
s.cell(row=yr_start, column=1, value="By Year").font = Font(name=F, size=11, bold=True, color="1F3A5F")
header(yr_start + 1, ["Year", "Sales", "Profit", "Margin", "YoY Sales Growth"])
years = sorted(df["Year"].unique())
for i, y in enumerate(years):
    r = yr_start + 2 + i
    put(r, 1, int(y)).number_format = "0"
    put(r, 2, f"=SUMIFS({rng('J')},{rng('L')},A{r})", USD)
    put(r, 3, f"=SUMIFS({rng('K')},{rng('L')},A{r})", USD)
    put(r, 4, f"=IF(B{r}=0,0,C{r}/B{r})", PCT)
    put(r, 5, "n/a" if i == 0 else f"=B{r}/B{r-1}-1", PCT)
yr_rows = (yr_start + 2, yr_start + 1 + len(years))

# Discount table
d_start = yr_rows[1] + 3
s.cell(row=d_start, column=1, value="By Discount Level").font = Font(name=F, size=11, bold=True, color="1F3A5F")
header(d_start + 1, ["Discount", "Orders", "Sales", "Profit", "Margin"])
discs = sorted(df["Discount"].unique())
for i, d in enumerate(discs):
    r = d_start + 2 + i
    put(r, 1, float(d), "0%")
    put(r, 2, f"=COUNTIF({rng('I')},A{r})", INT)
    put(r, 3, f"=SUMIFS({rng('J')},{rng('I')},A{r})", USD)
    put(r, 4, f"=SUMIFS({rng('K')},{rng('I')},A{r})", USD)
    put(r, 5, f"=IF(C{r}=0,0,D{r}/C{r})", PCT)
d_rows = (d_start + 2, d_start + 1 + len(discs))

for col, w in zip("ABCDE", [24, 16, 16, 12, 16]): s.column_dimensions[col].width = w

# Charts
def bar(title, rows, anchor, cols=(2, 3), horiz=False, width=15, height=7.5):
    ch = BarChart(); ch.type = "bar" if horiz else "col"; ch.title = title; ch.width, ch.height = width, height
    for c in cols:
        ch.add_data(Reference(s, min_col=c, min_row=rows[0] - 1, max_row=rows[1]), titles_from_data=True)
    ch.set_categories(Reference(s, min_col=1, min_row=rows[0], max_row=rows[1]))
    ch.y_axis.numFmt = '$#,##0'; ch.y_axis.delete = False; ch.x_axis.delete = False
    s.add_chart(ch, anchor)

bar("Sales and Profit by Category", cat_rows, "G4")
bar("Profit by Sub-Category", sub_rows, "G20", cols=(3,), horiz=True, height=9)
bar("Sales and Profit by Region", reg_rows, "G39")
lc = LineChart(); lc.title = "Yearly Sales and Profit"; lc.width, lc.height = 15, 7.5
for c in (2, 3):
    lc.add_data(Reference(s, min_col=c, min_row=yr_rows[0] - 1, max_row=yr_rows[1]), titles_from_data=True)
lc.set_categories(Reference(s, min_col=1, min_row=yr_rows[0], max_row=yr_rows[1]))
lc.y_axis.numFmt = '$#,##0'; lc.y_axis.delete = False; lc.x_axis.delete = False
s.add_chart(lc, "G55")
bar("Profit by Discount Level", d_rows, "G71", cols=(4,))
s.freeze_panes = "A4"

# ---------------- Model_Results ----------------
m = wb.create_sheet("Model_Results")
m["A1"] = "Regression Model Results (from Python notebook)"; m["A1"].font = TITLE
m["A2"] = ("Blue values are inputs copied from reports/metrics.json produced by notebooks/capstone_sales_analysis.ipynb. "
           "Target: Sales (log1p-transformed), 80/20 train-test split, random_state=42.")
m["A2"].font = Font(name=F, size=9, italic=True)
heads = ["Model", "R2 (test, USD)", "MAE (USD)", "RMSE (USD)", "CV R2 (log, 5-fold)"]
for j, h in enumerate(heads, 1):
    c = m.cell(row=4, column=j, value=h); c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
    c.alignment = Alignment(horizontal="center", wrap_text=True)
for i, r in enumerate(metrics["model_comparison"]):
    row = 5 + i
    vals = [r["Model"], r["R2 (test, USD)"], r["MAE (USD)"], r["RMSE (USD)"], r["CV R2 (log, 5-fold)"]]
    fmts = [None, "0.000", '$#,##0.00', '$#,##0.00', "0.000"]
    for j, (v, f) in enumerate(zip(vals, fmts), 1):
        c = m.cell(row=row, column=j, value=v); c.font = BODY if j == 1 else INPUT; c.border = BORDER
        if f: c.number_format = f
m["A9"] = "Best model (highest R2)"; m["A9"].font = BOLD
m["B9"] = "=INDEX(A5:A7,MATCH(MAX(B5:B7),B5:B7,0))"; m["B9"].font = BOLD
m["A11"] = "Feature importance (Gradient Boosting)"; m["A11"].font = Font(name=F, size=11, bold=True, color="1F3A5F")
for j, h in enumerate(["Feature", "Importance"], 1):
    c = m.cell(row=12, column=j, value=h); c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
for i, (k, v) in enumerate(metrics["feature_importance"].items()):
    m.cell(row=13 + i, column=1, value=k).font = BODY
    c = m.cell(row=13 + i, column=2, value=v); c.font = INPUT; c.number_format = "0.0%"
fi_end = 12 + len(metrics["feature_importance"])
ch = BarChart(); ch.type = "bar"; ch.title = "Feature Importance"; ch.width, ch.height = 14, 8
ch.add_data(Reference(m, min_col=2, min_row=12, max_row=fi_end), titles_from_data=True)
ch.set_categories(Reference(m, min_col=1, min_row=13, max_row=fi_end)); ch.legend = None
ch.y_axis.delete = False; ch.x_axis.delete = False
m.add_chart(ch, "D11")
for col, w in zip("ABCDE", [30, 16, 14, 14, 18]): m.column_dimensions[col].width = w
for row in m.iter_rows(min_row=1, max_row=3):
    for c in row: c.alignment = Alignment(wrap_text=False)

# ---------------- Notes ----------------
n = wb.create_sheet("Notes")
notes = [
    ("Project", "Week 4 Capstone: Retail Sales Analysis and Prediction"),
    ("Data source", "Synthetic retail dataset generated by src/generate_data.py (seed 42). The capstone brief supplied no dataset."),
    ("Cleaning", f"{metrics['duplicates_removed']} duplicate rows removed; missing Segment and Ship Mode filled with the mode."),
    ("Rows", f"{N:,} cleaned order lines, 2022-01-01 to 2025-12-31."),
    ("Formulas", "Year, Profit Margin and every Summary table are formulas. Model_Results blue cells are pasted inputs from Python."),
    ("Re-use", "To analyse real data, paste it into Clean_Data (same columns A:K); the Summary sheet recalculates. Update the row range if the row count changes."),
]
n["A1"] = "Notes and Assumptions"; n["A1"].font = TITLE
for i, (k, v) in enumerate(notes):
    n.cell(row=3 + i, column=1, value=k).font = BOLD
    c = n.cell(row=3 + i, column=2, value=v); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top")
n.column_dimensions["A"].width = 16; n.column_dimensions["B"].width = 110

out = ROOT / "excel" / "Retail_Sales_Analysis.xlsx"
wb.save(out); print("saved", out)
